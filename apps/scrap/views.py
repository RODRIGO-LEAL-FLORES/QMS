from functools import wraps
from decimal import Decimal,InvalidOperation
from datetime import datetime,date
from io import BytesIO
import logging
from urllib.parse import urlencode

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.mail import EmailMessage
from django.core.exceptions import ValidationError
from django.db import IntegrityError
from django.db.models import ProtectedError,Sum
from django.http import JsonResponse,HttpResponse
from django.shortcuts import render,redirect,get_object_or_404
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_POST
from django.core.paginator import Paginator

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.platypus import SimpleDocTemplate,Table,TableStyle,Paragraph,PageBreak,Flowable

from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter

from .models import Scrap,Operador,Turno,DefectoScrap,ClasificacionScrap,Supervisor,TipoAcero,TipoPieza
from apps.liberaciones.models import Maquina,TipoLaminacion
from apps.clientes.models import Cliente
from apps.usuarios.models import Usuario

logger=logging.getLogger(__name__)
LIMITE_PNC_KG=500

# ============================================================
# SEGURIDAD
# ============================================================
def nombre_rol(user):
    if not user.is_authenticated:return ''
    rol=getattr(user,'rol',None)
    if not rol:return ''
    return (getattr(rol,'nombre','') or '').strip().upper()

def puede_acceder_scrap(user):
    return bool(user.is_authenticated and getattr(user,'puede_gestionar_scrap',False))

def puede_administrar_scrap(user):
    return puede_acceder_scrap(user) and nombre_rol(user) in {'ADMIN','ENCARGADO'}

def requiere_gestion_scrap(view_func):
    @wraps(view_func)
    @login_required
    def wrapper(request,*args,**kwargs):
        if not puede_acceder_scrap(request.user):
            messages.error(request,'No tienes permiso para acceder al módulo de Scrap.')
            return redirect('home')
        return view_func(request,*args,**kwargs)
    return wrapper

def requiere_admin_scrap(view_func):
    @wraps(view_func)
    @login_required
    def wrapper(request,*args,**kwargs):
        if not puede_acceder_scrap(request.user):
            messages.error(request,'No tienes permiso para acceder al módulo de Scrap.')
            return redirect('home')
        if not puede_administrar_scrap(request.user):
            messages.error(request,'No tienes autorización para acceder a esta sección de Scrap.')
            return redirect('scrap_reportes')
        return view_func(request,*args,**kwargs)
    return wrapper


@requiere_gestion_scrap
def scrap_reporte_excel(request):
    if not request.user.is_superuser and request.user.rol_id not in (1,2):
        messages.error(request,'No tienes permiso para exportar a Excel.')
        return redirect('scrap_reportes')

# ============================================================
# UTILIDADES
# ============================================================
def decimal_post(request,nombre,default='0'):
    valor=(request.POST.get(nombre) or default).strip()
    try:return Decimal(valor)
    except (InvalidOperation,TypeError,ValueError):return Decimal(default)

def detectar_turno(hora_obj):
    for turno in Turno.objects.all():
        if turno.hora_inicio<=turno.hora_fin:
            if turno.hora_inicio<=hora_obj<turno.hora_fin:return turno
        else:
            if hora_obj>=turno.hora_inicio or hora_obj<turno.hora_fin:return turno
    return None

def construir_fecha_registro(fecha_str=None,hora_str=None,fecha_actual=None):
    ahora=timezone.localtime()
    if fecha_str and hora_str:
        dt=datetime.strptime(f'{fecha_str} {hora_str}','%Y-%m-%d %H:%M')
    elif fecha_str:
        hora_base=timezone.localtime(fecha_actual).time() if fecha_actual else ahora.time()
        dt=datetime.combine(datetime.strptime(fecha_str,'%Y-%m-%d').date(),hora_base.replace(tzinfo=None))
    elif hora_str:
        fecha_base=timezone.localtime(fecha_actual).date() if fecha_actual else ahora.date()
        dt=datetime.combine(fecha_base,datetime.strptime(hora_str,'%H:%M').time())
    else:return ahora
    if timezone.is_naive(dt):dt=timezone.make_aware(dt,timezone.get_current_timezone())
    return dt

def texto_clasificacion(clasificacion):
    return (getattr(clasificacion,'clasificacion','') or '').strip().lower()

def calcular_pnc(request,clasificacion):
    control=decimal_post(request,'control_salida_pnc_kg')
    agranel=decimal_post(request,'pnc_agranel_kg')
    nombre=texto_clasificacion(clasificacion)
    if nombre in ('targeta roja','tarjeta roja'):
        agranel=Decimal('0')
    elif nombre in ('bin amarillo','bin rojo'):
        control=Decimal('0')
    else:
        control=Decimal('0')
        agranel=Decimal('0')
    if control<0 or agranel<0:raise ValidationError('Los valores de PNC no pueden ser negativos.')
    return control,agranel,control+agranel

def mensaje_validacion(e):
    if hasattr(e,'message_dict'):
        salida=[]
        for errores in e.message_dict.values():salida.extend(errores)
        return ' '.join(salida)
    if hasattr(e,'messages'):return ' '.join(e.messages)
    return str(e)

# ============================================================
# ALERTA PNC
# ============================================================
def scrap_ticket_notification(registro):
    if not registro.pnc_kg or registro.pnc_kg<=LIMITE_PNC_KG:return
    usuarios=Usuario.objects.filter(
        puede_gestionar_scrap=True,
        is_active=True
    ).exclude(email__isnull=True).exclude(email='').distinct()
    destinatarios=list(usuarios.values_list('email',flat=True))
    if not destinatarios:return
    maquina=registro.maquina.nombre if registro.maquina else 'N/A'
    operador=registro.operador.nombre if registro.operador else 'N/A'
    defecto=registro.defecto.defecto if registro.defecto else 'N/A'
    cliente=registro.cliente.nombre if registro.cliente else 'N/A'
    fecha=timezone.localtime(registro.fecha_registro).strftime('%Y-%m-%d %H:%M') if registro.fecha_registro else 'N/A'
    asunto=f'Alerta de Scrap — Registro #{registro.id} excede el límite de PNC'
    cuerpo=f'''Hola,

Un registro de Scrap ha sobrepasado el límite de {LIMITE_PNC_KG} KG de PNC.

Registro: #{registro.id}
Máquina: {maquina}
Operador: {operador}
Defecto: {defecto}
Cliente: {cliente}
Número de orden: {registro.numero_de_orden or "N/A"}
Número de control: {registro.numero_de_control or "N/A"}
Lote: {registro.lote or "N/A"}
KG producidos: {registro.kg_producidos or 0}
PNC total: {registro.pnc_kg or 0} kg
Fecha: {fecha}

Por favor ingresa al sistema para revisar este registro.'''
    EmailMessage(subject=asunto,body=cuerpo,to=destinatarios).send(fail_silently=True)

# ============================================================
# MENÚ
# ============================================================
@login_required
def scrap(request):
    es_produccion=acceso_produccion_scrap(request.user) and not request.user.is_superuser and request.user.rol_id not in (1,2)
    return render(request,'scrap/scrap.html',{
        'es_produccion':es_produccion,
        'puede_administrar_scrap':puede_administrar_scrap(request.user),
        'puede_ver_produccion_scrap':acceso_produccion_scrap(request.user)
    })

# ============================================================
# CONTROLADOR PRINCIPAL
# ============================================================
@requiere_admin_scrap
def scrap_section(request,section):
    edit_id=request.GET.get('edit_id')

    if section=='nuevo':
        if request.method=='POST':
            try:
                fecha_manual=(request.POST.get('fecha_registro') or '').strip()
                hora_manual=(request.POST.get('hora_registro') or '').strip()
                fecha_final=construir_fecha_registro(fecha_manual or None,hora_manual or None)
                turno=detectar_turno(timezone.localtime(fecha_final).time().replace(tzinfo=None))
                if not turno:raise ValidationError('No existe un turno configurado para esta hora.')

                maquina=get_object_or_404(Maquina,id_maquina=request.POST.get('id_maquina'))
                operador=get_object_or_404(Operador,pk=request.POST.get('id_operador'))
                defecto=get_object_or_404(DefectoScrap,id_defecto_scrap=request.POST.get('id_defecto_scrap'))
                laminacion=get_object_or_404(TipoLaminacion,id_tipo_laminacion=request.POST.get('id_tipo_laminacion'))
                clasificacion=get_object_or_404(ClasificacionScrap,id_clasificacion=request.POST.get('id_clasificacion'))
                supervisor=get_object_or_404(Supervisor,id_supervisor=request.POST.get('id_supervisor'))
                cliente=get_object_or_404(Cliente,id_cliente=request.POST.get('id_cliente'))
                tipo_acero=get_object_or_404(TipoAcero,id_tipo_acero=request.POST.get('id_tipo_acero'))
                tipo_pieza=get_object_or_404(TipoPieza,id_tipo_pieza=request.POST.get('id_tipo_pieza'))

                numero_de_orden=(request.POST.get('numero_de_orden') or '').strip()
                numero_de_control=(request.POST.get('numero_de_control') or '').strip() or None
                lote=(request.POST.get('lote') or '').strip()
                contenedor=(request.POST.get('contenedor') or '').strip() or None
                kg_producidos=decimal_post(request,'kg_producidos')

                if not numero_de_orden:raise ValidationError('El número de orden es obligatorio.')
                if not lote:raise ValidationError('El lote es obligatorio.')
                if kg_producidos<0:raise ValidationError('Los KG producidos no pueden ser negativos.')

                control,agranel,pnc_total=calcular_pnc(request,clasificacion)

                nuevo_registro=Scrap(
                    fecha_registro=fecha_final,
                    maquina=maquina,
                    operador=operador,
                    turno=turno,
                    defecto=defecto,
                    clasificacion=clasificacion,
                    supervisor=supervisor,
                    cliente=cliente,
                    tipo_acero=tipo_acero,
                    usuario_registro=request.user,
                    tipo_laminacion=laminacion,
                    tipo_pieza=tipo_pieza,
                    numero_de_orden=numero_de_orden,
                    numero_de_control=numero_de_control,
                    lote=lote,
                    contenedor=contenedor,
                    control_salida_pnc_kg=control,
                    pnc_agranel_kg=agranel,
                    pnc_kg=pnc_total,
                    kg_producidos=kg_producidos
                )
                nuevo_registro.full_clean()
                nuevo_registro.save()

                try:scrap_ticket_notification(nuevo_registro)
                except Exception as e:logger.exception('Error enviando alerta PNC: %s',e)

                messages.success(request,'Registro de Scrap generado exitosamente.')
                return redirect('scrap_section',section='nuevo')
            except ValidationError as e:
                messages.error(request,mensaje_validacion(e))
            except Exception as e:
                logger.exception('Error guardando registro de Scrap')
                messages.error(request,f'Error al guardar: {e}')
            return redirect('scrap_section',section='nuevo')

        registros_qs=Scrap.objects.select_related(
            'maquina','operador','turno','defecto','tipo_laminacion',
            'clasificacion','supervisor','cliente','tipo_acero',
            'tipo_pieza','usuario_registro'
        ).order_by('-id')

        paginator=Paginator(registros_qs,10)
        page_obj=paginator.get_page(request.GET.get('page'))
        total_registros=registros_qs.count()
        kg_producidos_total=sum((r.kg_producidos or Decimal('0')) for r in registros_qs)
        pnc_total=sum((r.pnc_kg or Decimal('0')) for r in registros_qs)

        return render(request,'scrap/generar_registro.html',{
            'registros':page_obj.object_list,
            'pagination':page_obj,
            'page_obj':page_obj,
            'total_registros':total_registros,
            'kg_producidos_total':kg_producidos_total,
            'pnc_total':pnc_total,
            'peso_total':kg_producidos_total,
            'peso_ng_total':pnc_total,
            'maquinas':Maquina.objects.all().order_by('nombre'),
            'operadores':Operador.objects.all().order_by('nombre'),
            'turnos':Turno.objects.all().order_by('nombre_turno'),
            'defectos':DefectoScrap.objects.all().order_by('defecto'),
            'clasificaciones':ClasificacionScrap.objects.all().order_by('clasificacion'),
            'supervisores':Supervisor.objects.all().order_by('nombre'),
            'clientes':Cliente.objects.all().order_by('nombre'),
            'tipos_acero':TipoAcero.objects.all().order_by('especificacion'),
            'tipos_laminacion':TipoLaminacion.objects.all().order_by('especificacion'),
            'tipos_pieza':TipoPieza.objects.all().order_by('tipo_pieza')
        })

    mapping={
        'maquinas':(Maquina,'scrap/maquinas.html','nombre','id_maquina'),
        'operadores':(Operador,'scrap/operadores.html','nombre','pk'),
        'turnos':(Turno,'scrap/turnos.html','nombre_turno','id_turno'),
        'defectos':(DefectoScrap,'scrap/scrap_defectos.html','defecto','id_defecto_scrap'),
        'clasificaciones':(ClasificacionScrap,'scrap/clasificaciones.html','clasificacion','id_clasificacion'),
        'supervisores':(Supervisor,'scrap/supervisores.html','nombre','id_supervisor'),
        'clientes':(Cliente,'scrap/clientes_scrap.html','nombre','id_cliente'),
        'tipos_acero':(TipoAcero,'scrap/tipos_acero.html','especificacion','id_tipo_acero'),
        'tipos_laminacion':(TipoLaminacion,'scrap/tipos_laminacion.html','especificacion','id_tipo_laminacion'),
        'tipos_pieza':(TipoPieza,'scrap/tipos_pieza.html','tipo_pieza','id_tipo_pieza')
    }

    if section in mapping:
        model,template,field,pk_name=mapping[section]
        items=model.objects.all().order_by(field)
        edit_item=None
        if edit_id:
            try:edit_item=model.objects.filter(**{pk_name:edit_id}).first()
            except (ValueError,TypeError):edit_item=None
        return render(request,template,{'items':items,'edit_item':edit_item})

    messages.error(request,'La sección solicitada no existe.')
    return redirect('scrap')

# ============================================================
# CRUD CATÁLOGOS
# ============================================================
@requiere_admin_scrap
@require_POST
def scrap_actions(request,section,action_type,item_id=None):
    if section=='turnos':
        try:
            if action_type=='crear':
                nombre=(request.POST.get('nombre_turno') or '').strip()
                inicio=request.POST.get('hora_inicio')
                fin=request.POST.get('hora_fin')
                if not nombre or not inicio or not fin:raise ValidationError('Completa todos los campos.')
                if Turno.objects.filter(nombre_turno__iexact=nombre).exists():raise ValidationError('Este turno ya existe.')
                Turno.objects.create(
                    nombre_turno=nombre,
                    hora_inicio=datetime.strptime(inicio,'%H:%M').time(),
                    hora_fin=datetime.strptime(fin,'%H:%M').time()
                )
                messages.success(request,'Turno guardado con éxito.')
            elif action_type=='editar' and item_id:
                obj=get_object_or_404(Turno,id_turno=item_id)
                nombre=(request.POST.get('nombre_turno') or '').strip()
                if Turno.objects.filter(nombre_turno__iexact=nombre).exclude(id_turno=item_id).exists():
                    raise ValidationError('Ya existe otro turno con ese nombre.')
                obj.nombre_turno=nombre
                obj.hora_inicio=datetime.strptime(request.POST.get('hora_inicio'),'%H:%M').time()
                obj.hora_fin=datetime.strptime(request.POST.get('hora_fin'),'%H:%M').time()
                obj.save()
                messages.success(request,'Turno actualizado con éxito.')
            elif action_type=='eliminar' and item_id:
                get_object_or_404(Turno,id_turno=item_id).delete()
                messages.success(request,'Turno eliminado correctamente.')
            else:
                messages.error(request,'Acción no válida.')
        except ProtectedError:
            messages.error(request,'No se puede eliminar porque este turno está siendo utilizado.')
        except ValidationError as e:
            messages.error(request,mensaje_validacion(e))
        except Exception as e:
            messages.error(request,f'Error: {e}')
        return redirect('scrap_section',section='turnos')

    if section=='operadores':
        try:
            numero=(request.POST.get('numero_operador') or '').strip()
            nombre=(request.POST.get('nombre') or '').strip()
            if action_type in ('crear','editar') and not nombre:
                raise ValidationError('El nombre del operador es obligatorio.')

            if action_type=='crear':
                if numero and Operador.objects.filter(numero_operador__iexact=numero).exists():
                    raise ValidationError('Ese número de operador ya existe.')
                if Operador.objects.filter(nombre__iexact=nombre).exists():
                    raise ValidationError('Ese nombre de operador ya existe.')
                Operador.objects.create(numero_operador=numero or None,nombre=nombre)
                messages.success(request,'Operador creado correctamente.')

            elif action_type=='editar' and item_id:
                obj=get_object_or_404(Operador,pk=item_id)
                if numero and Operador.objects.filter(numero_operador__iexact=numero).exclude(pk=item_id).exists():
                    raise ValidationError('Ese número de operador ya existe.')
                if Operador.objects.filter(nombre__iexact=nombre).exclude(pk=item_id).exists():
                    raise ValidationError('Ese nombre de operador ya existe.')
                obj.numero_operador=numero or None
                obj.nombre=nombre
                obj.save()
                messages.success(request,'Operador actualizado correctamente.')

            elif action_type=='eliminar' and item_id:
                get_object_or_404(Operador,pk=item_id).delete()
                messages.success(request,'Operador eliminado correctamente.')
            else:
                messages.error(request,'Acción no válida.')
        except ProtectedError:
            messages.error(request,'No se puede eliminar porque este operador está siendo utilizado.')
        except ValidationError as e:
            messages.error(request,mensaje_validacion(e))
        except Exception as e:
            messages.error(request,f'Error: {e}')
        return redirect('scrap_section',section='operadores')

    if section=='tipos_pieza':
        try:
            nombre=(request.POST.get('tipo_pieza') or '').strip()
            costo=decimal_post(request,'costo_unitario')

            if action_type in ('crear','editar') and not nombre:
                raise ValidationError('El tipo de pieza es obligatorio.')
            if costo<0:raise ValidationError('El costo unitario no puede ser negativo.')

            if action_type=='crear':
                if TipoPieza.objects.filter(tipo_pieza__iexact=nombre).exists():
                    raise ValidationError('Este tipo de pieza ya existe.')
                TipoPieza.objects.create(tipo_pieza=nombre,costo_unitario=costo)
                messages.success(request,'Tipo de pieza creado correctamente.')

            elif action_type=='editar' and item_id:
                obj=get_object_or_404(TipoPieza,id_tipo_pieza=item_id)
                if TipoPieza.objects.filter(tipo_pieza__iexact=nombre).exclude(id_tipo_pieza=item_id).exists():
                    raise ValidationError('Ya existe otro tipo de pieza con ese nombre.')
                obj.tipo_pieza=nombre
                obj.costo_unitario=costo
                obj.save()
                messages.success(request,'Tipo de pieza actualizado correctamente.')

            elif action_type=='eliminar' and item_id:
                get_object_or_404(TipoPieza,id_tipo_pieza=item_id).delete()
                messages.success(request,'Tipo de pieza eliminado correctamente.')
            else:
                messages.error(request,'Acción no válida.')
        except ProtectedError:
            messages.error(request,'No se puede eliminar porque este tipo de pieza está siendo utilizado.')
        except ValidationError as e:
            messages.error(request,mensaje_validacion(e))
        except Exception as e:
            messages.error(request,f'Error: {e}')
        return redirect('scrap_section',section='tipos_pieza')

    model_mapping={
        'maquinas':(Maquina,'nombre','id_maquina'),
        'defectos':(DefectoScrap,'defecto','id_defecto_scrap'),
        'clasificaciones':(ClasificacionScrap,'clasificacion','id_clasificacion'),
        'supervisores':(Supervisor,'nombre','id_supervisor'),
        'tipos_acero':(TipoAcero,'especificacion','id_tipo_acero'),
        'clientes':(Cliente,'nombre','id_cliente'),
        'tipos_laminacion':(TipoLaminacion,'especificacion','id_tipo_laminacion')
    }

    if section not in model_mapping:
        messages.error(request,'Sección no válida.')
        return redirect('scrap')

    model,field_name,pk_name=model_mapping[section]

    try:
        if action_type in ('crear','editar'):
            value=(request.POST.get(field_name) or '').strip()
            if not value:raise ValidationError('El campo requerido no puede estar vacío.')

            if action_type=='crear':
                if model.objects.filter(**{f'{field_name}__iexact':value}).exists():
                    raise ValidationError('Este registro ya existe en el sistema.')
                datos={field_name:value}
                if section=='maquinas':
                    datos['descripcion']=(request.POST.get('descripcion') or '').strip() or None
                model.objects.create(**datos)
                messages.success(request,'Registro creado con éxito.')
            else:
                obj=get_object_or_404(model,**{pk_name:item_id})
                if model.objects.filter(**{f'{field_name}__iexact':value}).exclude(**{pk_name:item_id}).exists():
                    raise ValidationError('Ya existe otro registro con ese mismo nombre.')
                setattr(obj,field_name,value)
                if section=='maquinas':
                    obj.descripcion=(request.POST.get('descripcion') or '').strip() or None
                obj.save()
                messages.success(request,'Registro actualizado con éxito.')

        elif action_type=='eliminar' and item_id:
            get_object_or_404(model,**{pk_name:item_id}).delete()
            messages.success(request,'Registro eliminado correctamente.')
        else:
            messages.error(request,'Acción no válida.')

    except ProtectedError:
        messages.error(request,'No se puede eliminar este registro porque está siendo utilizado.')
    except IntegrityError:
        messages.error(request,'No se pudo realizar la operación porque el registro está relacionado con otros datos.')
    except ValidationError as e:
        messages.error(request,mensaje_validacion(e))
    except Exception as e:
        messages.error(request,f'Error: {e}')

    return redirect('scrap_section',section=section)

# ============================================================
# EDITAR / ELIMINAR SCRAP
# ============================================================
@requiere_admin_scrap
@require_POST
def scrap_editar(request,item_id):
    registro=get_object_or_404(Scrap,id=item_id)

    try:
        registro.maquina=get_object_or_404(Maquina,id_maquina=request.POST.get('id_maquina'))
        registro.operador=get_object_or_404(Operador,pk=request.POST.get('id_operador'))
        registro.supervisor=get_object_or_404(Supervisor,id_supervisor=request.POST.get('id_supervisor'))
        registro.defecto=get_object_or_404(DefectoScrap,id_defecto_scrap=request.POST.get('id_defecto_scrap'))
        registro.clasificacion=get_object_or_404(ClasificacionScrap,id_clasificacion=request.POST.get('id_clasificacion'))
        registro.cliente=get_object_or_404(Cliente,id_cliente=request.POST.get('id_cliente'))
        registro.tipo_acero=get_object_or_404(TipoAcero,id_tipo_acero=request.POST.get('id_tipo_acero'))
        registro.tipo_laminacion=get_object_or_404(TipoLaminacion,id_tipo_laminacion=request.POST.get('id_tipo_laminacion'))
        registro.tipo_pieza=get_object_or_404(TipoPieza,id_tipo_pieza=request.POST.get('id_tipo_pieza'))

        registro.numero_de_orden=(request.POST.get('numero_de_orden') or '').strip()
        registro.numero_de_control=(request.POST.get('numero_de_control') or '').strip() or None
        registro.lote=(request.POST.get('lote') or '').strip()
        registro.contenedor=(request.POST.get('contenedor') or '').strip() or None
        registro.kg_producidos=decimal_post(request,'kg_producidos')

        if not registro.numero_de_orden:raise ValidationError('El número de orden es obligatorio.')
        if not registro.lote:raise ValidationError('El lote es obligatorio.')
        if registro.kg_producidos<0:raise ValidationError('Los KG producidos no pueden ser negativos.')

        control,agranel,pnc_total=calcular_pnc(request,registro.clasificacion)
        registro.control_salida_pnc_kg=control
        registro.pnc_agranel_kg=agranel
        registro.pnc_kg=pnc_total

        hora_manual=(request.POST.get('hora_registro') or '').strip()
        fecha_manual=(request.POST.get('fecha_registro') or '').strip()

        if hora_manual or fecha_manual:
            fecha_final=construir_fecha_registro(
                fecha_manual or None,
                hora_manual or None,
                registro.fecha_registro
            )
            turno=detectar_turno(timezone.localtime(fecha_final).time().replace(tzinfo=None))
            if not turno:raise ValidationError('No existe un turno configurado para esta hora.')
            registro.fecha_registro=fecha_final
            registro.turno=turno

        registro.full_clean()
        registro.save()

        try:scrap_ticket_notification(registro)
        except Exception as e:logger.exception('Error enviando alerta PNC: %s',e)

        messages.success(request,'Registro actualizado correctamente.')

    except ValidationError as e:
        messages.error(request,mensaje_validacion(e))
    except Exception as e:
        logger.exception('Error actualizando Scrap #%s',item_id)
        messages.error(request,f'Error al actualizar: {e}')

    return redirect('scrap_section',section='nuevo')

@requiere_admin_scrap
@require_POST
def scrap_eliminar(request,item_id):
    registro=get_object_or_404(Scrap,id=item_id)
    try:
        registro.delete()
        messages.success(request,'Registro eliminado correctamente.')
    except ProtectedError:
        messages.error(request,'No se puede eliminar porque el registro está relacionado con otros datos.')
    except Exception as e:
        messages.error(request,f'Error al eliminar: {e}')
    return redirect('scrap_section',section='nuevo')

# ============================================================
# HISTORIAL / TURNO
# ============================================================
@requiere_admin_scrap
def historial_usuario(request,user_id):
    registros=Scrap.objects.filter(usuario_registro_id=user_id).select_related(
        'maquina','operador','turno','defecto','cliente','tipo_acero',
        'tipo_laminacion','tipo_pieza','clasificacion','supervisor'
    ).order_by('-fecha_registro')
    return render(request,'scrap/historial.html',{'registros':registros})

@requiere_admin_scrap
def get_turno(request):
    hora_str=request.GET.get('hora')
    if not hora_str:return JsonResponse({'error':'Hora no proporcionada'},status=400)
    try:
        turno=detectar_turno(datetime.strptime(hora_str,'%H:%M').time())
        if turno:return JsonResponse({'id_turno':turno.id_turno,'nombre':turno.nombre_turno})
        return JsonResponse({'id_turno':None,'nombre':'Sin turno asignado'})
    except Exception as e:
        return JsonResponse({'error':str(e)},status=500)

# ============================================================
# FILTROS
# ============================================================
def obtener_filtros_scrap(request):
    return {
        'fecha_inicio':request.GET.get('fecha_inicio',''),
        'fecha_fin':request.GET.get('fecha_fin',''),
        'id_maquina':request.GET.get('id_maquina',''),
        'id_operador':request.GET.get('id_operador',''),
        'id_cliente':request.GET.get('id_cliente',''),
        'id_defecto_scrap':request.GET.get('id_defecto_scrap',''),
        'id_turno':request.GET.get('id_turno',''),
        'id_supervisor':request.GET.get('id_supervisor',''),
        'id_clasificacion':request.GET.get('id_clasificacion',''),
        'id_tipo_acero':request.GET.get('id_tipo_acero',''),
        'id_tipo_laminacion':request.GET.get('id_tipo_laminacion',''),
        'id_tipo_pieza':request.GET.get('id_tipo_pieza','')
    }

def aplicar_filtros_scrap(query,filtros):
    if filtros['fecha_inicio']:
        inicio=timezone.make_aware(
            datetime.strptime(filtros['fecha_inicio'],'%Y-%m-%d'),
            timezone.get_current_timezone()
        )
        query=query.filter(fecha_registro__gte=inicio)

    if filtros['fecha_fin']:
        fin=datetime.strptime(filtros['fecha_fin'],'%Y-%m-%d').replace(hour=23,minute=59,second=59)
        fin=timezone.make_aware(fin,timezone.get_current_timezone())
        query=query.filter(fecha_registro__lte=fin)

    mapping={
        'id_maquina':'maquina_id',
        'id_operador':'operador_id',
        'id_cliente':'cliente_id',
        'id_defecto_scrap':'defecto_id',
        'id_turno':'turno_id',
        'id_supervisor':'supervisor_id',
        'id_clasificacion':'clasificacion_id',
        'id_tipo_acero':'tipo_acero_id',
        'id_tipo_laminacion':'tipo_laminacion_id',
        'id_tipo_pieza':'tipo_pieza_id'
    }

    for filtro,campo in mapping.items():
        if filtros[filtro]:
            query=query.filter(**{campo:filtros[filtro]})

    return query

def queryset_reporte_scrap(request):
    filtros=obtener_filtros_scrap(request)
    query=Scrap.objects.select_related(
        'maquina','operador','turno','defecto','clasificacion',
        'supervisor','cliente','tipo_acero','usuario_registro',
        'tipo_laminacion','tipo_pieza'
    )
    query=aplicar_filtros_scrap(query,filtros)
    return query,filtros

# ============================================================
# REPORTES
# ============================================================
@requiere_gestion_scrap
def scrap_reportes(request):
    filtros=obtener_filtros_scrap(request)
    query=Scrap.objects.select_related(
        'maquina','operador','turno','defecto','clasificacion',
        'supervisor','cliente','tipo_acero','usuario_registro',
        'tipo_laminacion','tipo_pieza'
    )

    try:
        query=aplicar_filtros_scrap(query,filtros)
    except (ValueError,TypeError):
        messages.error(request,'Uno de los filtros contiene un valor inválido.')

    registros=query.order_by('-fecha_registro')
    total_kg=sum((r.kg_producidos or Decimal('0')) for r in registros)
    total_pnc=sum((r.pnc_kg or Decimal('0')) for r in registros)
    total_control=sum((r.control_salida_pnc_kg or Decimal('0')) for r in registros)
    total_agranel=sum((r.pnc_agranel_kg or Decimal('0')) for r in registros)
    count=registros.count()

    kpis={
        'total_registros':count,
        'kg_producidos':total_kg,
        'pnc_total':total_pnc,
        'control_salida_pnc':total_control,
        'pnc_agranel':total_agranel,
        'pnc_promedio':(total_pnc/count if count else Decimal('0')),
        'peso_total':total_kg,
        'total_ng':total_pnc,
        'total_retrabajo':total_agranel,
        'peso_promedio':(total_kg/count if count else Decimal('0'))
    }

    return render(request,'scrap/reportes_scrap.html',{
        'filtros':filtros,
        'registros':registros,
        'kpis':kpis,
        'maquinas':Maquina.objects.all().order_by('nombre'),
        'operadores':Operador.objects.all().order_by('nombre'),
        'clientes':Cliente.objects.all().order_by('nombre'),
        'defectos':DefectoScrap.objects.all().order_by('defecto'),
        'turnos':Turno.objects.all().order_by('nombre_turno'),
        'supervisores':Supervisor.objects.all().order_by('nombre'),
        'clasificaciones':ClasificacionScrap.objects.all().order_by('clasificacion'),
        'tipos_acero':TipoAcero.objects.all().order_by('especificacion'),
        'tipos_laminacion':TipoLaminacion.objects.all().order_by('especificacion'),
        'tipos_pieza':TipoPieza.objects.all().order_by('tipo_pieza')
    })

# ============================================================
# FORMATO OPERATIVO PDF
# ============================================================
@requiere_gestion_scrap
def scrap_formato_pdf(request):
    class Checkbox(Flowable):
        def __init__(self,size=4*mm):
            super().__init__()
            self.width=size
            self.height=size
            self.size=size

        def draw(self):
            self.canv.setStrokeColor(colors.black)
            self.canv.setLineWidth(.8)
            self.canv.rect(0,0,self.size,self.size,stroke=1,fill=0)

    try:
        query,filtros=queryset_reporte_scrap(request)
    except (ValueError,TypeError):
        messages.error(request,'Los filtros proporcionados no son válidos.')
        return redirect('scrap_reportes')

    registros=list(query.order_by('fecha_registro'))
    buffer=BytesIO()
    doc=SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=7*mm,
        leftMargin=7*mm,
        topMargin=8*mm,
        bottomMargin=10*mm
    )

    styles=getSampleStyleSheet()
    s_t=ParagraphStyle(
        'titulo_scrap',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=13,
        spaceAfter=4
    )
    s_n=ParagraphStyle(
        'normal_scrap',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=6,
        leading=7
    )
    s_h=ParagraphStyle(
        'header_scrap',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=4.7,
        leading=5.1,
        alignment=1
    )
    s_c=ParagraphStyle(
        'celda_scrap',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=5.1,
        leading=5.6,
        alignment=1
    )

    story=[]
    grupos={}

    for r in registros:
        fecha=timezone.localtime(r.fecha_registro).date() if r.fecha_registro else None
        grupos.setdefault(fecha,[]).append(r)

    if not grupos:
        grupos[timezone.localdate()]=[]

    for gi,(fecha,regs) in enumerate(grupos.items()):
        bloques=[regs[i:i+18] for i in range(0,len(regs),18)] or [[]]

        for bi,bloque in enumerate(bloques):
            if gi>0 or bi>0:
                story.append(PageBreak())

            story.append(Paragraph('Registro de scrap',s_t))
            ft=fecha.strftime('%d/%m/%Y') if fecha else ''

            tf=Table(
                [[Paragraph('<b>FECHA</b>',s_h),Paragraph(ft,s_c)]],
                colWidths=[20*mm,176*mm],
                rowHeights=[8*mm]
            )
            tf.setStyle(TableStyle([
                ('GRID',(0,0),(-1,-1),.5,colors.black),
                ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
                ('ALIGN',(0,0),(-1,-1),'CENTER')
            ]))
            story.append(tf)

            filas=[[
                Paragraph('HORA',s_h),
                Paragraph('# DE<br/>INSPECTOR',s_h),
                Paragraph('MÁQUINA',s_h),
                Paragraph('LAMINACIÓN',s_h),
                Paragraph('DEFECTO',s_h),
                Paragraph('PNC KG',s_h),
                Paragraph('RETRABAJO',s_h),
                Paragraph('ORDEN DE<br/>TRABAJO<br/>ASIGNADA',s_h),
                Paragraph('POSTEADO<br/>EN EL ERP',s_h),
                Paragraph('# DE<br/>OPERADOR',s_h),
                Paragraph('# DE<br/>CONTROL',s_h)
            ]]

            for r in bloque:
                hora=timezone.localtime(r.fecha_registro).strftime('%H:%M') if r.fecha_registro else ''
                inspector=getattr(r.usuario_registro,'username','') if r.usuario_registro else ''

                filas.append([
                    Paragraph(hora,s_c),
                    Paragraph(str(inspector),s_c),
                    Paragraph(r.maquina.nombre if r.maquina else '',s_c),
                    Paragraph(r.tipo_laminacion.especificacion if r.tipo_laminacion else '',s_c),
                    Paragraph(r.defecto.defecto if r.defecto else '',s_c),
                    Paragraph(f'{float(r.pnc_kg or 0):,.1f}',s_c),
                    Checkbox(),
                    Paragraph(r.numero_de_orden or '',s_c),
                    Checkbox(),
                    Paragraph(str(getattr(r.operador,'numero_operador','') or '') if r.operador else '',s_c),
                    Paragraph(r.numero_de_control or '',s_c)
                ])

            while len(filas)<19:
                filas.append(['','','','','','',Checkbox(),'',Checkbox(),'',''])

            anchos=[
                13*mm,18*mm,20*mm,24*mm,25*mm,
                13*mm,17*mm,22*mm,18*mm,14*mm,12*mm
            ]

            tabla=Table(
                filas,
                colWidths=anchos,
                rowHeights=[10*mm]+[8*mm]*18,
                repeatRows=1
            )
            tabla.setStyle(TableStyle([
                ('GRID',(0,0),(-1,-1),.45,colors.black),
                ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#f2f2f2')),
                ('VALIGN',(0,0),(-1,-1),'MIDDLE'),
                ('ALIGN',(0,0),(-1,-1),'CENTER'),
                ('LEFTPADDING',(0,0),(-1,-1),1),
                ('RIGHTPADDING',(0,0),(-1,-1),1)
            ]))
            story.append(tabla)

            val=Table([
                [
                    Paragraph('<b>VALIDA</b>',s_h),
                    Paragraph('<b>VALIDA</b>',s_h),
                    Paragraph('<b>LIBERA</b>',s_h)
                ],
                ['','','']
            ],colWidths=[65.33*mm]*3,rowHeights=[7*mm,14*mm])

            val.setStyle(TableStyle([
                ('GRID',(0,0),(-1,-1),.45,colors.black),
                ('ALIGN',(0,0),(-1,-1),'CENTER'),
                ('VALIGN',(0,0),(-1,-1),'MIDDLE')
            ]))
            story.append(val)

            firmas=Table([[
                Paragraph('Firma del supervisor 1er. turno',s_n),
                Paragraph('Firma del supervisor 2do. turno',s_n),
                Paragraph('Firma del supervisor de calidad',s_n)
            ]],colWidths=[65.33*mm]*3,rowHeights=[10*mm])

            firmas.setStyle(TableStyle([
                ('GRID',(0,0),(-1,-1),.45,colors.black),
                ('ALIGN',(0,0),(-1,-1),'CENTER'),
                ('VALIGN',(0,0),(-1,-1),'MIDDLE')
            ]))
            story.append(firmas)

    def footer(canvas,doc):
        canvas.saveState()
        ancho,_=A4
        canvas.setFont('Helvetica-Bold',7)
        canvas.drawString(7*mm,5*mm,'VC LAMINATIONS')
        canvas.setFont('Helvetica',5.5)
        canvas.drawRightString(ancho-7*mm,5*mm,f'Página {doc.page}')
        canvas.restoreState()

    doc.build(story,onFirstPage=footer,onLaterPages=footer)
    buffer.seek(0)

    nombre=f'registro_scrap_{timezone.localtime().strftime("%Y%m%d_%H%M")}.pdf'
    response=HttpResponse(buffer.getvalue(),content_type='application/pdf')
    response['Content-Disposition']=f'attachment; filename="{nombre}"'
    return response

# ============================================================
# EXPORTAR EXCEL
# ============================================================
@requiere_gestion_scrap
def scrap_reporte_excel(request):
    try:
        query,filtros=queryset_reporte_scrap(request)
    except (ValueError,TypeError):
        messages.error(request,'Los filtros proporcionados no son válidos.')
        return redirect('scrap_reportes')

    registros=query.order_by('fecha_registro')

    wb=Workbook()
    ws=wb.active
    ws.title='Scrap'

    encabezados=[
        'DÍA',
        'MES',
        'SEMANA',
        'PERIODO',
        'CONTENEDOR',
        'TURNO',
        'PRENSA',
        'LAMINACIÓN',
        'NÚMERO DE CONTROL',
        'CLIENTE',
        'PNC CONTROLES DE SALIDA DE SCRAP',
        'PNC A GRANEL',
        'KG PNC',
        'KG PRODUCIDOS',
        'DEFECTO',
        'TIPO',
        'COSTO'
    ]

    for columna,encabezado in enumerate(encabezados,1):
        celda=ws.cell(row=1,column=columna,value=encabezado)
        celda.font=Font(bold=True,color='FFFFFF')
        celda.fill=PatternFill('solid',fgColor='1F4E78')
        celda.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)

    borde=Border(
        left=Side(style='thin',color='BFBFBF'),
        right=Side(style='thin',color='BFBFBF'),
        top=Side(style='thin',color='BFBFBF'),
        bottom=Side(style='thin',color='BFBFBF')
    )

    fila=2

    for registro in registros:
        fecha=timezone.localtime(registro.fecha_registro) if registro.fecha_registro else None
        iso=fecha.isocalendar() if fecha else None
        costo=registro.tipo_pieza.costo_unitario if registro.tipo_pieza else Decimal('0')

        valores=[
            fecha.day if fecha else '',
            fecha.month if fecha else '',
            iso.week if iso else '',
            fecha.strftime('%m-%Y') if fecha else '',
            registro.contenedor or '',
            registro.turno.nombre_turno if registro.turno else '',
            registro.maquina.nombre if registro.maquina else '',
            registro.tipo_laminacion.especificacion if registro.tipo_laminacion else '',
            registro.numero_de_control or '',
            registro.cliente.nombre if registro.cliente else '',
            float(registro.control_salida_pnc_kg or 0),
            float(registro.pnc_agranel_kg or 0),
            float(registro.pnc_kg or 0),
            float(registro.kg_producidos or 0),
            registro.defecto.defecto if registro.defecto else '',
            registro.tipo_pieza.tipo_pieza if registro.tipo_pieza else '',
            float(costo or 0)
        ]

        for columna,valor in enumerate(valores,1):
            celda=ws.cell(row=fila,column=columna,value=valor)
            celda.border=borde
            celda.alignment=Alignment(vertical='center',wrap_text=True)

        for columna in range(11,15):
            ws.cell(row=fila,column=columna).number_format='#,##0.00'

        ws.cell(row=fila,column=17).number_format='$#,##0.00'
        fila+=1

    for columna in range(1,len(encabezados)+1):
        ws.cell(row=1,column=columna).border=borde

    anchos={
        'A':10,'B':10,'C':10,'D':12,
        'E':18,'F':15,'G':18,'H':25,
        'I':22,'J':25,'K':30,'L':20,
        'M':15,'N':18,'O':30,'P':20,'Q':15
    }

    for columna,ancho in anchos.items():
        ws.column_dimensions[columna].width=ancho

    ws.row_dimensions[1].height=45
    ws.freeze_panes='A2'
    ws.auto_filter.ref=f'A1:{get_column_letter(len(encabezados))}{max(fila-1,1)}'

    buffer=BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre=f'scrap_{timezone.localtime().strftime("%Y%m%d_%H%M")}.xlsx'
    response=HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition']=f'attachment; filename="{nombre}"'
    return response


# ============================================================
# PRODUCCIÓN: CONSULTA POR DÍA Y EDICIÓN LIMITADA
# ============================================================
def acceso_produccion_scrap(user):
    if not user.is_authenticated or not user.is_active:return False
    if user.is_superuser:return True
    area=getattr(user,'area',None)
    nombre=(getattr(area,'nombre','') or '').strip().casefold()
    return nombre in ('producción','produccion')


@login_required
def produccion_scrap(request):
    if not acceso_produccion_scrap(request.user):
        messages.error(request,'No tienes permiso para acceder a Producción.')
        return redirect('home')
    fecha=request.GET.get('fecha') or timezone.localdate().isoformat()
    turno_id=request.GET.get('turno','').strip()
    orden=request.GET.get('orden','').strip()
    trabajador=request.GET.get('trabajador','').strip()
    try:
        fecha_seleccionada=date.fromisoformat(fecha)
    except ValueError:
        fecha_seleccionada=timezone.localdate()
    registros=Scrap.objects.filter(
        fecha_registro__date=fecha_seleccionada
    ).select_related(
        'maquina','operador','turno','cliente','defecto',
        'clasificacion','tipo_laminacion','tipo_pieza'
    )
    if turno_id:
        if turno_id.isdigit() and Turno.objects.filter(pk=int(turno_id)).exists():
            registros=registros.filter(turno_id=int(turno_id))
        else:
            turno_id=''
    if orden:
        registros=registros.filter(numero_de_orden__icontains=orden)
    if trabajador:
        registros=registros.filter(
            Q(operador__nombre__icontains=trabajador) |
            Q(operador__numero_operador__icontains=trabajador)
        )
    registros=registros.order_by('-fecha_registro','-id')
    totales=registros.aggregate(
        kg_producidos=Sum('kg_producidos'),
        kg_pnc=Sum('pnc_kg')
    )
    return render(request,'scrap/produccion_scrap.html',{
        'registros':registros,
        'turnos':Turno.objects.all().order_by('id_turno'),
        'turno_seleccionado':turno_id,
        'fecha_seleccionada':fecha_seleccionada.isoformat(),
        'orden_busqueda':orden,
        'trabajador_busqueda':trabajador,
        'total_registros':registros.count(),
        'total_kg':totales['kg_producidos'] or Decimal('0'),
        'total_pnc':totales['kg_pnc'] or Decimal('0')
    })



@login_required
@require_POST
def produccion_scrap_editar(request,item_id):
    if not acceso_produccion_scrap(request.user):
        messages.error(request,'No tienes permiso para editar registros de Producción.')
        return redirect('home')

    registro=get_object_or_404(Scrap,pk=item_id)

    fecha=timezone.localtime(registro.fecha_registro).date().isoformat()
    turno=(request.POST.get('turno') or '').strip()
    orden=(request.POST.get('orden') or '').strip()
    trabajador=(request.POST.get('trabajador') or '').strip()

    try:
        control=(request.POST.get('numero_de_control') or '').strip() or None
        kg_texto=(request.POST.get('kg_producidos') or '').strip()

        if control and len(control)>50:
            raise ValidationError('El número de control admite máximo 50 caracteres.')

        if not kg_texto:
            raise ValidationError('Los KG producidos son obligatorios.')

        try:
            kg=Decimal(kg_texto)
        except (InvalidOperation,ValueError):
            raise ValidationError('Introduce un valor válido de KG producidos.')

        if not kg.is_finite() or kg<0 or kg>=Decimal('100000000'):
            raise ValidationError('Los KG producidos deben ser un número no negativo de máximo 8 dígitos enteros.')

        if kg.as_tuple().exponent<-2:
            raise ValidationError('Los KG producidos admiten máximo dos decimales.')

        registro.numero_de_control=control
        registro.kg_producidos=kg

        registro.full_clean()
        registro.save(update_fields=['numero_de_control','kg_producidos'])

        messages.success(
            request,
            f'Registro #{registro.id} actualizado correctamente.'
        )

    except ValidationError as e:
        messages.error(request,mensaje_validacion(e))

    except Exception:
        logger.exception(
            'Error actualizando campos de Producción en Scrap #%s',
            item_id
        )
        messages.error(request,'No fue posible actualizar el registro.')

    parametros={'fecha':fecha}

    if turno.isdigit() and Turno.objects.filter(pk=int(turno)).exists():
        parametros['turno']=turno

    if orden:
        parametros['orden']=orden

    if trabajador:
        parametros['trabajador']=trabajador

    return redirect(
        reverse('produccion_scrap')+'?'+urlencode(parametros)
    )